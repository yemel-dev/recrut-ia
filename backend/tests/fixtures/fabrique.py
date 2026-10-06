"""Fabrique de CV de test : PDF texte, PDF « scanné » (image seule) et DOCX, à partir des textes de fixtures/cv/."""
from __future__ import annotations

import zlib
from pathlib import Path

DOSSIER_CV = Path(__file__).parent / "cv"


def texte_cv(nom: str) -> str:
    return (DOSSIER_CV / f"{nom}.txt").read_text(encoding="utf-8")


def _echapper(ligne: str) -> bytes:
    brut = ligne.encode("cp1252", errors="replace")
    return brut.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _assembler_pdf(objets: list[bytes]) -> bytes:
    """objets[i] est le contenu de l'objet i+1 (sans « n 0 obj »)."""
    sortie = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    positions = []
    for numero, contenu in enumerate(objets, start=1):
        positions.append(len(sortie))
        sortie += f"{numero} 0 obj\n".encode() + contenu + b"\nendobj\n"
    debut_xref = len(sortie)
    sortie += f"xref\n0 {len(objets) + 1}\n0000000000 65535 f \n".encode()
    for position in positions:
        sortie += f"{position:010d} 00000 n \n".encode()
    sortie += f"trailer\n<< /Size {len(objets) + 1} /Root 1 0 R >>\nstartxref\n{debut_xref}\n%%EOF\n".encode()
    return bytes(sortie)


def _flux(donnees: bytes, dictionnaire: str = "") -> bytes:
    compresse = zlib.compress(donnees)
    return f"<< {dictionnaire} /Filter /FlateDecode /Length {len(compresse)} >>\nstream\n".encode() + compresse + b"\nendstream"


def pdf_texte(texte: str) -> bytes:
    """PDF d'une ou plusieurs pages A4, police Helvetica, encodage WinAnsi (accents compris)."""
    lignes_texte = texte.splitlines()
    pages = [lignes_texte[i : i + 60] for i in range(0, len(lignes_texte), 60)] or [[]]
    # 1 catalogue, 2 arbre des pages, 3 police, puis pour chaque page : page + contenu
    objets: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"]
    enfants = []
    for lignes_page in pages:
        contenu = bytearray(b"BT /F1 10 Tf 12 TL 50 800 Td\n")
        for ligne in lignes_page:
            contenu += b"(" + _echapper(ligne) + b") Tj T*\n"
        contenu += b"ET"
        numero_page = len(objets) + 1
        objets.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents {numero_page + 1} 0 R >>".encode())
        objets.append(_flux(bytes(contenu)))
        enfants.append(f"{numero_page} 0 R")
    objets[1] = f"<< /Type /Pages /Kids [{' '.join(enfants)}] /Count {len(enfants)} >>".encode()
    return _assembler_pdf(objets)


def pdf_scanne() -> bytes:
    """PDF d'une page qui ne contient qu'une image (comme un CV scanné) : aucun texte extractible."""
    largeur = hauteur = 64
    pixels = bytes((x * y) % 256 for y in range(hauteur) for x in range(largeur))
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /XObject << /Im1 4 0 R >> >> /Contents 5 0 R >>",
        _flux(pixels, f"/Type /XObject /Subtype /Image /Width {largeur} /Height {hauteur} /ColorSpace /DeviceGray /BitsPerComponent 8"),
        _flux(b"q 500 0 0 700 50 70 cm /Im1 Do Q"),
    ]
    return _assembler_pdf(objets)


# Polices système avec accents (celle par défaut de Pillow n'a pas « é » : elle dessine un carré vide).
POLICES = [
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]


def police_avec_accents() -> str | None:
    return next((p for p in POLICES if Path(p).exists()), None)


def pdf_scanne_texte(texte: str) -> bytes:
    """PDF d'une page A4 qui ne contient qu'une image du texte (CV scanné lisible), pour tester l'OCR."""
    from PIL import Image, ImageDraw, ImageFont

    largeur, hauteur = 1240, 1754  # A4 à 150 points par pouce
    image = Image.new("L", (largeur, hauteur), 255)
    dessin = ImageDraw.Draw(image)
    police = ImageFont.truetype(police_avec_accents(), 24)
    for i, ligne in enumerate(texte.splitlines()[:50]):
        dessin.text((80, 80 + i * 32), ligne, fill=0, font=police)
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /XObject << /Im1 4 0 R >> >> /Contents 5 0 R >>",
        _flux(image.tobytes(), f"/Type /XObject /Subtype /Image /Width {largeur} /Height {hauteur} /ColorSpace /DeviceGray /BitsPerComponent 8"),
        _flux(b"q 595 0 0 842 0 0 cm /Im1 Do Q"),
    ]
    return _assembler_pdf(objets)


def docx(texte: str, dans_un_tableau: bool = False) -> bytes:
    """DOCX ; avec dans_un_tableau=True, le corps du CV est mis en page dans un tableau à deux colonnes."""
    import io

    from docx import Document

    document = Document()
    lignes_texte = texte.splitlines()
    if dans_un_tableau:
        document.add_paragraph(lignes_texte[0])
        tableau = document.add_table(rows=1, cols=2)
        milieu = len(lignes_texte) // 2
        tableau.cell(0, 0).text = "\n".join(lignes_texte[1:milieu])
        tableau.cell(0, 1).text = "\n".join(lignes_texte[milieu:])
    else:
        for ligne in lignes_texte:
            document.add_paragraph(ligne)
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


def ecrire(dossier: Path, nom_fichier: str, contenu: bytes) -> Path:
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = dossier / nom_fichier
    chemin.write_bytes(contenu)
    return chemin
