"""Étape 1 : lecture du texte d'un CV (PDF ou DOCX).

Un fichier dont on ne tire pas assez de texte (scan, fichier vide, protégé ou corrompu) lève FichierIllisible :
la candidature est alors marquée « illisible » et signalée au recruteur, jamais notée 0 en silence.
"""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path

log = logging.getLogger("injara.ia.lecture")

# Nombre minimal de caractères utiles (lettres et chiffres) pour considérer un CV comme lisible.
LONGUEUR_MIN_TEXTE = 150

_UTILE = re.compile(r"[^\W_]", re.UNICODE)


class FichierIllisible(Exception):
    """Le texte du CV ne peut pas être lu. `motif` est destiné au recruteur."""

    def __init__(self, motif: str) -> None:
        super().__init__(motif)
        self.motif = motif


def lire_texte(chemin: str | Path, contenu: bytes | None = None) -> str:
    """Texte brut du CV. Lève FichierIllisible si le fichier est inexploitable.

    `contenu` : octets déjà lus (CV déchiffré en mémoire) ; `chemin` ne sert alors qu'au nom et à l'extension.
    """
    chemin = Path(chemin)
    if contenu is None:
        if not chemin.exists():
            raise FichierIllisible("Fichier introuvable : il a peut-être été déplacé ou supprimé.")
        contenu = chemin.read_bytes()
    extension = chemin.suffix.lower()
    if extension == ".pdf":
        texte = _lire_pdf(contenu, chemin.name)
    elif extension == ".docx":
        texte = _lire_docx(contenu, chemin.name)
    else:
        raise FichierIllisible(f"Format non pris en charge ({extension or 'sans extension'}) : PDF ou DOCX attendu.")

    utiles = len(_UTILE.findall(texte))
    if utiles < LONGUEUR_MIN_TEXTE:
        if extension == ".pdf":
            raise FichierIllisible(
                "Aucun texte lisible dans ce PDF : il s'agit probablement d'un document scanné (image). "
                "Ouvrez-le pour le consulter."
            )
        raise FichierIllisible("Ce document ne contient presque pas de texte.")
    log.info("Lecture : %d caractères utiles (%s)", utiles, chemin.name)
    return texte.strip()


def _lire_pdf(contenu: bytes, nom: str) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover - dépendance du socle
        raise FichierIllisible("Lecteur PDF absent (pypdf) : installez requirements.txt.") from exc
    try:
        lecteur = PdfReader(io.BytesIO(contenu))
        if lecteur.is_encrypted and not lecteur.decrypt(""):
            raise FichierIllisible("Ce PDF est protégé par un mot de passe.")
        if not lecteur.pages:
            raise FichierIllisible("Ce PDF ne contient aucune page.")
        return "\n".join(page.extract_text() or "" for page in lecteur.pages)
    except FichierIllisible:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        log.warning("PDF illisible %s : %s", nom, exc)
        raise FichierIllisible("Ce PDF est endommagé ou dans un format non reconnu.") from exc


def _lire_docx(contenu: bytes, nom: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dépendance du socle
        raise FichierIllisible("Lecteur DOCX absent (python-docx) : installez requirements.txt.") from exc
    try:
        document = Document(io.BytesIO(contenu))
    except Exception as exc:  # zip invalide, XML corrompu…
        log.warning("DOCX illisible %s : %s", nom, exc)
        raise FichierIllisible("Ce document Word est endommagé ou dans un format non reconnu.") from exc
    morceaux = [p.text for p in document.paragraphs]
    # Beaucoup de CV Word sont mis en page dans des tableaux : on lit aussi leurs cellules, ligne par ligne.
    for tableau in document.tables:
        for rangee in tableau.rows:
            vues: set[int] = set()
            for cellule in rangee.cells:
                if id(cellule._tc) in vues:  # cellules fusionnées : une seule lecture
                    continue
                vues.add(id(cellule._tc))
                morceaux.append(cellule.text)
    return "\n".join(m for m in morceaux if m.strip())
